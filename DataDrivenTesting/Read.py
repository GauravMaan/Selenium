import os

import openpyxl
workbook=openpyxl.load_workbook('/Users/gauravmaan/Downloads/Test case (open cart).xlsx')

sheet=workbook["Home"]
row=sheet.max_row
print(row)
col=sheet.max_column
print(col)
for r in range(1,row+1):
    for c in range(1,col+1):
        print(sheet.cell(r,c).value,end="                              ")
    print()
