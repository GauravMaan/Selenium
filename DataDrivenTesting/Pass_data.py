# import time
# import openpyxl
# from selenium import webdriver
# from selenium.webdriver.common.by import By
# from selenium.webdriver.chrome.options import Options
# from selenium.webdriver.support.ui import WebDriverWait
# from selenium.webdriver.support import expected_conditions as EC
#
# file = '/Users/gauravmaan/Downloads/simple_interest_data.xlsx'
# workbook = openpyxl.load_workbook(file)
# sheet = workbook.active
# deposit_amount = sheet.cell(2, 1).value
#
# interest_rate = sheet.cell(2, 2).value
# tenure = sheet.cell(2, 3).value
#
# chrome_options = Options()
# driver = webdriver.Chrome(options=chrome_options)
# driver.get("https://www.moneycontrol.com/fixed-income/calculator/state-bank-of-india-sbi/fixed-deposit-calculator-SBI-BSB001.html")
#
# wait = WebDriverWait(driver, 10)
# deposit_field = wait.until(EC.presence_of_element_located((By.XPATH, "//*[@id='principal']")))
# deposit_field.clear()
# deposit_field.send_keys(str(deposit_amount))
#
# time.sleep(2)
# rate_field = driver.find_element(By.ID, "interest")
# rate_field.clear()
# rate_field.send_keys(str(interest_rate))
#
# time.sleep(2)
# tenure_field = driver.find_element(By.ID, "tenure")
# tenure_field.clear()
# tenure_field.send_keys(str(tenure))
#
# time.sleep(2)
# calculate_button = driver.find_element(By.XPATH, "//*[@id='fdMatVal']/div[2]/a[1]/img")
# calculate_button.click()
#
# time.sleep(3)
# driver.quit()
#
import time
import openpyxl
from selenium import webdriver
from selenium.webdriver.common.by import By
from selenium.webdriver.chrome.options import Options
from selenium.webdriver.support.ui import WebDriverWait
from selenium.webdriver.support import expected_conditions as EC

file = '/Users/gauravmaan/Downloads/simple_interest_data.xlsx'
workbook = openpyxl.load_workbook(file)
sheet = workbook.active

chrome_options = Options()
driver = webdriver.Chrome(options=chrome_options)
driver.get("https://www.moneycontrol.com/fixed-income/calculator/state-bank-of-india-sbi/fixed-deposit-calculator-SBI-BSB001.html")

wait = WebDriverWait(driver, 10)

for row in range(2, sheet.max_row + 1):
    deposit_amount = sheet.cell(row, 1).value
    interest_rate = sheet.cell(row, 2).value
    tenure = sheet.cell(row, 3).value

    deposit_field = wait.until(EC.presence_of_element_located((By.XPATH, "//*[@id='principal']")))
    deposit_field.clear()
    deposit_field.send_keys(str(deposit_amount))

    time.sleep(2)
    rate_field = driver.find_element(By.ID, "interest")
    rate_field.clear()
    rate_field.send_keys(str(interest_rate))

    time.sleep(2)
    tenure_field = driver.find_element(By.ID, "tenure")
    tenure_field.clear()
    tenure_field.send_keys(str(tenure))

    time.sleep(2)
    calculate_button = driver.find_element(By.XPATH, "//*[@id='fdMatVal']/div[2]/a[1]/img")
    calculate_button.click()

    time.sleep(3)

driver.quit()

